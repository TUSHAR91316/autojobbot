import re
from pathlib import Path
from typing import Dict, Any, List
from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract plain text from PDF using pypdf"""
    text_chunks = []
    try:
        reader = PdfReader(str(file_path))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")
    return "\n".join(text_chunks).strip()


def extract_text_from_docx(file_path: Path) -> str:
    """Extract plain text from DOCX using python-docx"""
    text_chunks = []
    try:
        doc = Document(str(file_path))
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_chunks.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_chunks.append(cell.text)
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
    return "\n".join(text_chunks).strip()


def parse_resume_content(file_path: Path, file_type: str) -> Dict[str, Any]:
    """Extract text and detect basic contact details safely without hallucination"""
    if file_type.lower() == "pdf":
        text = extract_text_from_pdf(file_path)
    elif file_type.lower() in ["docx", "doc"]:
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
    
    # Safe regex extractors
    email_match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text)
    phone_match = re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
    
    detected_email = email_match.group(0) if email_match else ""
    detected_phone = phone_match.group(0) if phone_match else ""
    
    return {
        "text": text,
        "detected_email": detected_email,
        "detected_phone": detected_phone,
    }
